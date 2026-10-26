"""
Production-Ready Hybrid Identity Document Extractor (Streamlit Web Application)
===============================================================================
Extracts core Aadhaar & Indian ID fields (12-digit ID Number, Holder Name,
Date of Birth, Gender, Father/Guardian Name, Mobile Number, and Address)
using a Hybrid Qwen2-VL-2B (Google Drive / Local) + EasyOCR + PyMuPDF + Regex
pipeline across single and multi-page documents.
"""

import collections
import importlib.util
import io
import json
import os
import re
import shutil
import sys
import tempfile
import time
import urllib.parse
from typing import Any, Dict, List, Optional, Tuple

import cv2
import numpy as np
import pandas as pd
from PIL import Image
import streamlit as st

# Ensure UTF-8 output across Windows consoles
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

# Fast availability checks without eager heavy ML imports on startup
EASYOCR_AVAILABLE = importlib.util.find_spec("easyocr") is not None
TORCH_AVAILABLE = importlib.util.find_spec("torch") is not None
QWEN_VL_AVAILABLE = (
    importlib.util.find_spec("transformers") is not None
    and importlib.util.find_spec("qwen_vl_utils") is not None
)
HF_HUB_AVAILABLE = importlib.util.find_spec("huggingface_hub") is not None

try:
    import fitz  # PyMuPDF
    PYMUPDF_AVAILABLE = True
except ImportError:
    PYMUPDF_AVAILABLE = False

try:
    import pdf2image
    PDF2IMAGE_AVAILABLE = True
except ImportError:
    PDF2IMAGE_AVAILABLE = False

try:
    import gdown
    GDOWN_AVAILABLE = True
except ImportError:
    GDOWN_AVAILABLE = False

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


# ==============================================================================
# QWEN2-VL-2B GOOGLE DRIVE CHECKPOINT VALIDATOR & MODEL LOADER
# ==============================================================================

DEFAULT_QWEN_CANDIDATE_PATHS = [
    "/content/drive/MyDrive/model_ocr/qwen2_vl_2b",
    "/content/drive/MyDrive/qwen2_vl_2b",
    os.path.expanduser("~/Google Drive/My Drive/model_ocr/qwen2_vl_2b"),
    "G:/My Drive/model_ocr/qwen2_vl_2b",
]
DEFAULT_QWEN_HUB_ID = "Qwen/Qwen2-VL-2B-Instruct"


def verify_qwen_checkpoint(model_dir: str) -> Tuple[bool, List[str]]:
    """
    Verifies that a local or Google Drive Qwen2-VL checkpoint directory is complete
    and that all required .safetensors (or .bin) weight shards exist and are non-empty.
    Prevents FileNotFoundError during Qwen2VLForConditionalGeneration.from_pretrained().
    Returns (is_valid, list_of_missing_or_corrupt_files).
    """
    if not model_dir or not os.path.isdir(model_dir):
        return False, ["directory_not_found"]

    missing: List[str] = []
    config_path = os.path.join(model_dir, "config.json")
    if not os.path.isfile(config_path) or os.path.getsize(config_path) == 0:
        missing.append("config.json")

    safetensors_index = os.path.join(model_dir, "model.safetensors.index.json")
    bin_index = os.path.join(model_dir, "pytorch_model.bin.index.json")

    shard_files: List[str] = []
    if os.path.isfile(safetensors_index) and os.path.getsize(safetensors_index) > 0:
        try:
            with open(safetensors_index, "r", encoding="utf-8") as f:
                idx_data = json.load(f)
            weight_map = idx_data.get("weight_map", {})
            shard_files = sorted(set(weight_map.values()))
        except Exception:
            missing.append("model.safetensors.index.json (corrupt)")
    elif os.path.isfile(bin_index) and os.path.getsize(bin_index) > 0:
        try:
            with open(bin_index, "r", encoding="utf-8") as f:
                idx_data = json.load(f)
            weight_map = idx_data.get("weight_map", {})
            shard_files = sorted(set(weight_map.values()))
        except Exception:
            missing.append("pytorch_model.bin.index.json (corrupt)")
    else:
        # Check single-file weights or standard 2-shard Qwen2-VL-2B layout
        single_st = os.path.join(model_dir, "model.safetensors")
        single_bin = os.path.join(model_dir, "pytorch_model.bin")
        if os.path.isfile(single_st) and os.path.getsize(single_st) > 0:
            shard_files = ["model.safetensors"]
        elif os.path.isfile(single_bin) and os.path.getsize(single_bin) > 0:
            shard_files = ["pytorch_model.bin"]
        else:
            shard_files = [
                "model-00001-of-00002.safetensors",
                "model-00002-of-00002.safetensors",
            ]

    for shard in shard_files:
        shard_path = os.path.join(model_dir, shard)
        if not os.path.isfile(shard_path) or os.path.getsize(shard_path) == 0:
            missing.append(shard)

    return len(missing) == 0, missing


def resolve_or_repair_qwen_model_path(
    candidate_paths: Optional[List[str]] = None,
    default_hub_id: str = DEFAULT_QWEN_HUB_ID,
    auto_repair: bool = True,
) -> str:
    """
    Scans candidate Google Drive / local directories for Qwen2-VL-2B.
    - If a complete checkpoint is found, returns its path immediately.
    - If an incomplete checkpoint is found (e.g. missing model-00001-of-00002.safetensors)
      and auto_repair=True, downloads missing files from Hugging Face Hub into that Drive folder.
    - Otherwise falls back cleanly to default_hub_id ('Qwen/Qwen2-VL-2B-Instruct').
    """
    paths = candidate_paths if candidate_paths is not None else DEFAULT_QWEN_CANDIDATE_PATHS

    for path in paths:
        if not path:
            continue
        if os.path.isdir(path):
            is_valid, missing_files = verify_qwen_checkpoint(path)
            if is_valid:
                return path

            print(
                f"⚠️ Incomplete Qwen2-VL checkpoint at '{path}'. "
                f"Missing/empty files: {missing_files}"
            )
            if auto_repair and HF_HUB_AVAILABLE:
                try:
                    from huggingface_hub import snapshot_download
                    print(f"🔄 Auto-repairing Drive checkpoint '{path}' from '{default_hub_id}'...")
                    snapshot_download(
                        repo_id=default_hub_id,
                        local_dir=path,
                        local_dir_use_symlinks=False,
                    )
                    re_valid, re_missing = verify_qwen_checkpoint(path)
                    if re_valid:
                        print(f"✅ Successfully repaired checkpoint at: {path}")
                        return path
                    print(f"⚠️ Checkpoint still incomplete after repair attempt: {re_missing}")
                except Exception as e:
                    print(f"⚠️ Auto-repair to Drive failed ({e}); falling back.")

    return default_hub_id


@st.cache_resource(show_spinner="Loading Qwen2-VL-2B Vision-Language Model...")
def load_qwen_vl_model(
    custom_path: Optional[str] = None,
    auto_repair: bool = True,
) -> Tuple[Optional[Any], Optional[Any], str]:
    """
    Loads and caches Qwen2-VL-2B model and AutoProcessor with hardware-appropriate dtype
    (bfloat16 on Ampere+ GPUs, float16 on Tesla T4, float32 on CPU).
    """
    if not QWEN_VL_AVAILABLE or not TORCH_AVAILABLE:
        return None, None, "Unavailable (Install transformers, qwen-vl-utils, accelerate)"

    import torch
    from transformers import Qwen2VLForConditionalGeneration, AutoProcessor

    candidates = [custom_path] + DEFAULT_QWEN_CANDIDATE_PATHS if custom_path else DEFAULT_QWEN_CANDIDATE_PATHS
    resolved_path = resolve_or_repair_qwen_model_path(
        candidate_paths=candidates,
        default_hub_id=DEFAULT_QWEN_HUB_ID,
        auto_repair=auto_repair,
    )

    if torch.cuda.is_available():
        compute_dtype = torch.bfloat16 if torch.cuda.is_bf16_supported() else torch.float16
        device_map = "auto"
    else:
        compute_dtype = torch.float32
        device_map = None

    model = Qwen2VLForConditionalGeneration.from_pretrained(
        resolved_path,
        torch_dtype=compute_dtype,
        device_map=device_map,
        low_cpu_mem_usage=True,
    )
    processor = AutoProcessor.from_pretrained(
        resolved_path,
        min_pixels=256 * 28 * 28,
        max_pixels=1280 * 28 * 28,
    )
    return model, processor, resolved_path


def extract_fields_with_qwen(
    pil_image: Image.Image,
    model: Any,
    processor: Any,
) -> Dict[str, Any]:
    """
    Runs Qwen2-VL-2B inference on a single document page image to extract Aadhaar / ID fields.
    """
    default_result = {
        "document_id": "Not Found",
        "holder_name": "Not Found",
        "date_of_birth": "Not Found",
        "gender": "Not Found",
        "father_or_guardian": "Not Found",
        "mobile_number": "Not Found",
        "address": "Not Found",
    }
    if model is None or processor is None or pil_image is None:
        return default_result

    import torch
    from qwen_vl_utils import process_vision_info

    vlm_prompt = (
        "You are an expert Indian Aadhaar and identity document OCR extractor.\n"
        "Analyze this document image (Front, Back, or full e-Aadhaar page) and extract these exact fields:\n"
        "- document_id: The 12-digit Aadhaar number (e.g. '1234 5678 9012' or 'XXXX XXXX 1234'). Return 'Not Found' if absent.\n"
        "- holder_name: Full name of the cardholder (usually above DOB/Gender on the front).\n"
        "- date_of_birth: Date of birth in DD/MM/YYYY or YYYY format.\n"
        "- gender: 'Male', 'Female', or 'Other'.\n"
        "- father_or_guardian: Name of the father, mother, husband, or guardian (look after S/O, D/O, W/O, C/O, Care of, or Father's Name, including inside the Address block).\n"
        "- mobile_number: 10-digit Indian mobile number starting with 6, 7, 8, or 9 if printed.\n"
        "- address: Full residential address including house/street/village/city/state and 6-digit PIN code if present.\n\n"
        "Output strictly valid JSON with these exact keys and no extra commentary:\n"
        "{\n"
        '  "document_id": "...",\n'
        '  "holder_name": "...",\n'
        '  "date_of_birth": "...",\n'
        '  "gender": "...",\n'
        '  "father_or_guardian": "...",\n'
        '  "mobile_number": "...",\n'
        '  "address": "..."\n'
        "}"
    )

    messages = [
        {
            "role": "user",
            "content": [
                {"type": "image", "image": pil_image},
                {"type": "text", "text": vlm_prompt},
            ],
        }
    ]

    try:
        text_prompt = processor.apply_chat_template(
            messages, tokenize=False, add_generation_prompt=True
        )
        image_inputs, _ = process_vision_info(messages)
        target_device = next(model.parameters()).device
        inputs = processor(
            text=[text_prompt],
            images=image_inputs,
            padding=True,
            return_tensors="pt",
        ).to(target_device)

        with torch.no_grad():
            generated_ids = model.generate(
                **inputs,
                max_new_tokens=384,
                do_sample=False,
                repetition_penalty=1.05,
            )

        in_len = inputs.input_ids.shape[-1]
        output_ids = generated_ids[:, in_len:]
        raw_response = processor.batch_decode(output_ids, skip_special_tokens=True)[0].strip()

        json_match = re.search(r"\{.*\}", raw_response, re.DOTALL)
        parsed = json.loads(json_match.group(0) if json_match else raw_response)
        merged = dict(default_result)
        if isinstance(parsed, dict):
            for k in default_result:
                val = parsed.get(k)
                if val is not None and str(val).strip() and str(val).strip().lower() not in {"none", "null", "n/a", "..."}:
                    merged[k] = str(val).strip()
        merged["raw_output"] = raw_response
        return merged
    except Exception as e:
        res = dict(default_result)
        res["raw_output"] = f"Qwen2-VL inference error: {e}"
        return res


# ==============================================================================
# MODEL CACHING & EASYOCR INITIALIZATION
# ==============================================================================

@st.cache_resource(show_spinner="Loading EasyOCR model weights into memory...")
def get_easyocr_reader(languages: Tuple[str, ...]) -> Optional[Any]:
    """Loads and caches EasyOCR reader instance."""
    if not EASYOCR_AVAILABLE:
        return None
    import easyocr
    gpu_available = False
    if TORCH_AVAILABLE:
        try:
            import torch
            gpu_available = torch.cuda.is_available()
        except Exception:
            gpu_available = False
    return easyocr.Reader(list(languages), gpu=gpu_available)


# ==============================================================================
# RESILIENT GOOGLE DRIVE INGESTION ENGINE
# ==============================================================================

def extract_gdrive_id(url_or_id: str) -> Optional[str]:
    """
    Robustly parses Google Drive URL schemes to extract the underlying File or Folder ID.
    Handles:
      - /file/d/<ID>/view?usp=sharing
      - /open?id=<ID>
      - /uc?id=<ID>
      - /folders/<ID>
      - Raw 25+ character ID strings
    """
    if not url_or_id:
        return None
    raw = url_or_id.strip()

    patterns = [
        r"/file/d/([a-zA-Z0-9_-]+)",
        r"[?&]id=([a-zA-Z0-9_-]+)",
        r"/folders/([a-zA-Z0-9_-]+)",
        r"/document/d/([a-zA-Z0-9_-]+)",
        r"/spreadsheets/d/([a-zA-Z0-9_-]+)",
        r"/presentation/d/([a-zA-Z0-9_-]+)",
        r"^([a-zA-Z0-9_-]{25,})$",
    ]
    for pattern in patterns:
        m = re.search(pattern, raw)
        if m:
            return m.group(1)
    return raw


def download_from_gdrive(url_or_id: str, dest_dir: str) -> List[str]:
    """
    Downloads document(s) from a Google Drive URL using gdown.
    Handles redirects, confirmation tokens, and fuzzy matching.
    Detects true file format using magic bytes.
    """
    if not GDOWN_AVAILABLE:
        raise RuntimeError("gdown library is required for Google Drive downloads.")

    item_id = extract_gdrive_id(url_or_id)
    is_folder = "folder" in url_or_id.lower()

    if is_folder and item_id:
        try:
            downloaded = gdown.download_folder(id=item_id, output=dest_dir, quiet=False, remaining_ok=True)
            if downloaded:
                valid_files = [f for f in downloaded if os.path.isfile(f) and os.path.getsize(f) > 0]
                if valid_files:
                    return valid_files
        except Exception as e:
            print(f"Folder download notice, attempting standard download: {e}")

    out_file = os.path.join(dest_dir, "gdrive_downloaded_document")
    try:
        gdown.download(id=item_id, output=out_file, quiet=False, fuzzy=True)
    except Exception:
        gdown.download(url=url_or_id.strip(), output=out_file, quiet=False, fuzzy=True)

    if not os.path.exists(out_file) or os.path.getsize(out_file) == 0:
        raise FileNotFoundError(
            "Could not download file from Google Drive. Please ensure the link is publicly accessible "
            "('Anyone with the link can view') or upload the file directly."
        )

    with open(out_file, "rb") as f:
        header = f.read(16)

    ext = ".bin"
    if header.startswith(b"%PDF"):
        ext = ".pdf"
    elif header.startswith(b"\xff\xd8\xff"):
        ext = ".jpg"
    elif header.startswith(b"\x89PNG"):
        ext = ".png"
    elif header.startswith(b"II*\x00") or header.startswith(b"MM\x00*"):
        ext = ".tiff"
    elif header.startswith(b"RIFF") and b"WEBP" in header:
        ext = ".webp"
    elif header.startswith(b"BM"):
        ext = ".bmp"

    final_path = out_file + ext
    shutil.move(out_file, final_path)
    return [final_path]


# ==============================================================================
# DIGITAL PDF EXTRACTION & NON-DESTRUCTIVE RASTERIZATION
# ==============================================================================

def extract_pdf_digital_text(pdf_path: str) -> Tuple[str, List[str]]:
    """
    Extracts text vectors directly from digitally issued e-PDFs using PyMuPDF.
    """
    if not PYMUPDF_AVAILABLE:
        return "", []

    all_lines = []
    try:
        doc = fitz.open(pdf_path)
        for page in doc:
            page_text = page.get_text("text")
            for line in page_text.splitlines():
                clean = line.strip()
                if clean:
                    all_lines.append(clean)
        doc.close()
    except Exception as e:
        print(f"Digital vector extraction notice: {e}")

    return "\n".join(all_lines), all_lines


def load_document_pages(file_path: str, display_name: str) -> List[Tuple[np.ndarray, str]]:
    """
    Loads all pages of an image or PDF into OpenCV BGR numpy arrays at 300 DPI.
    """
    pages = []
    is_pdf = file_path.lower().endswith(".pdf")
    if not is_pdf:
        try:
            with open(file_path, "rb") as f:
                header = f.read(8)
                if header.startswith(b"%PDF"):
                    is_pdf = True
        except Exception:
            pass

    if is_pdf:
        rendered = False
        if PYMUPDF_AVAILABLE:
            try:
                doc = fitz.open(file_path)
                for idx, page in enumerate(doc):
                    pix = page.get_pixmap(dpi=300)
                    img_bytes = pix.tobytes("png")
                    nparr = np.frombuffer(img_bytes, np.uint8)
                    cv_img = cv2.imdecode(nparr, cv2.IMREAD_COLOR)
                    pages.append((cv_img, f"{display_name} (Page {idx + 1})"))
                doc.close()
                rendered = True
            except Exception as e:
                print(f"PyMuPDF rasterization fallback to pdf2image: {e}")

        if not rendered and PDF2IMAGE_AVAILABLE:
            try:
                pil_imgs = pdf2image.convert_from_path(file_path, dpi=300)
                for idx, img in enumerate(pil_imgs):
                    cv_img = cv2.cvtColor(np.array(img), cv2.COLOR_RGB2BGR)
                    pages.append((cv_img, f"{display_name} (Page {idx + 1})"))
                rendered = True
            except Exception as e:
                raise RuntimeError(f"Failed to rasterize PDF: {e}")

        if not rendered:
            raise RuntimeError("Neither PyMuPDF nor pdf2image is available to render PDF pages.")
    else:
        with open(file_path, "rb") as f:
            file_bytes = f.read()
        nparr = np.frombuffer(file_bytes, np.uint8)
        cv_img = cv2.imdecode(nparr, cv2.IMREAD_COLOR)
        if cv_img is None:
            raise ValueError(f"Could not decode image at {file_path}")
        pages.append((cv_img, display_name))

    return pages


# ==============================================================================
# MULTI-PASS / NON-DESTRUCTIVE PREPROCESSING & OCR PIPELINE
# ==============================================================================

def sort_ocr_results(ocr_results: List[Any]) -> List[str]:
    """
    Sorts detected text bounding boxes in natural reading order without hardcoded slicing.
    """
    if not ocr_results:
        return []

    heights = [abs(b[2][1] - b[0][1]) for b, _, _ in ocr_results if abs(b[2][1] - b[0][1]) > 0]
    median_h = float(np.median(heights)) if heights else 20.0
    y_tolerance = max(10.0, median_h * 0.5)

    sorted_by_y = sorted(ocr_results, key=lambda item: min(p[1] for p in item[0]))

    lines: List[List[Any]] = []
    current_line: List[Any] = []
    current_y: Optional[float] = None

    for item in sorted_by_y:
        y_min = min(p[1] for p in item[0])
        if current_y is None or abs(y_min - current_y) <= y_tolerance:
            current_line.append(item)
            current_y = y_min if current_y is None else (current_y + y_min) / 2.0
        else:
            current_line.sort(key=lambda it: min(p[0] for p in it[0]))
            lines.append(current_line)
            current_line = [item]
            current_y = y_min

    if current_line:
        current_line.sort(key=lambda it: min(p[0] for p in it[0]))
        lines.append(current_line)

    flat_lines: List[str] = []
    for line_group in lines:
        text_str = " ".join(it[1].strip() for it in line_group if it[1].strip())
        if text_str:
            flat_lines.append(text_str)

    return flat_lines


def run_dual_pass_ocr(
    cv_img: np.ndarray,
    reader: Any,
    use_clahe_pass: bool = True,
) -> Tuple[str, List[str]]:
    """
    Executes an enhanced multi-pass OCR pipeline (Grayscale + Auto-Rotation + CLAHE).
    """
    if reader is None:
        return "", []

    if len(cv_img.shape) == 3:
        gray = cv2.cvtColor(cv_img, cv2.COLOR_BGR2GRAY)
    else:
        gray = cv_img.copy()

    h, w = gray.shape[:2]
    max_dim = max(h, w)
    if max_dim < 1600:
        scale = min(2.0, 1600.0 / max_dim)
        gray = cv2.resize(gray, None, fx=scale, fy=scale, interpolation=cv2.INTER_CUBIC)

    res1 = reader.readtext(gray, detail=1, paragraph=False)
    lines1 = sort_ocr_results(res1)
    all_lines = list(lines1)

    if len(all_lines) < 3:
        best_rot_lines = all_lines
        best_gray = gray
        for rot_code in [cv2.ROTATE_90_CLOCKWISE, cv2.ROTATE_180, cv2.ROTATE_90_COUNTERCLOCKWISE]:
            rot_img = cv2.rotate(gray, rot_code)
            res_rot = reader.readtext(rot_img, detail=1, paragraph=False)
            lines_rot = sort_ocr_results(res_rot)
            if len(lines_rot) > len(best_rot_lines):
                best_rot_lines = lines_rot
                best_gray = rot_img
        if len(best_rot_lines) > len(all_lines):
            gray = best_gray
            all_lines = list(best_rot_lines)

    if use_clahe_pass:
        clahe = cv2.createCLAHE(clipLimit=1.8, tileGridSize=(8, 8))
        enhanced = clahe.apply(gray)
        res2 = reader.readtext(enhanced, detail=1, paragraph=False)
        lines2 = sort_ocr_results(res2)

        existing_norm = {re.sub(r"\s+", "", l.lower()) for l in all_lines}
        for l in lines2:
            norm_l = re.sub(r"\s+", "", l.lower())
            if norm_l and norm_l not in existing_norm:
                all_lines.append(l)
                existing_norm.add(norm_l)

    full_corpus = "\n".join(all_lines)
    return full_corpus, all_lines


# ==============================================================================
# NOISE-TOLERANT CONTEXTUAL FIELD PARSERS & HYBRID FUSION
# ==============================================================================

def parse_document_id(text: str, lines: List[str]) -> Tuple[str, str]:
    """
    Noise-tolerant extraction of 12-Digit Indian Identity Number (starting with 2-9)
    or Masked Aadhaar (XXXX XXXX 1234 or **** **** 1234).
    Returns (masked_id, raw_id).
    """
    if not text and not lines:
        return "Not Found", "Not Found"

    masked_pattern = re.compile(r"\b([Xx*]{4})[\s\-\.\_]*([Xx*]{4})[\s\-\.\_]*(\d{4})\b")
    m_mask = masked_pattern.search(text)
    if m_mask:
        last4 = m_mask.group(3)
        return f"XXXX XXXX {last4}", f"XXXX XXXX {last4}"

    for line in lines:
        m_mask = masked_pattern.search(line)
        if m_mask:
            last4 = m_mask.group(3)
            return f"XXXX XXXX {last4}", f"XXXX XXXX {last4}"

    pattern_spaced = r"\b([2-9]\d{3})[\s\-\.\_]+(\d{4})[\s\-\.\_]+(\d{4})\b"
    match = re.search(pattern_spaced, text)
    if match:
        p1, p2, p3 = match.group(1), match.group(2), match.group(3)
        raw_id = f"{p1} {p2} {p3}"
        masked_id = f"XXXX XXXX {p3}"
        return masked_id, raw_id

    trans_table = str.maketrans({
        "O": "0", "o": "0", "Q": "0", "D": "0",
        "I": "1", "l": "1", "|": "1", "!": "1",
        "Z": "2", "z": "2",
        "S": "5", "s": "5",
        "B": "8",
    })

    for line in lines:
        if re.search(r"[xX*]{4}", line):
            m = re.search(r"(\d{4})\b", line)
            if m:
                return f"XXXX XXXX {m.group(1)}", f"XXXX XXXX {m.group(1)}"

        m_cand = re.search(
            r"\b([2-9A-Za-z][0-9A-Za-z]{3})[\s\-\.\_]+([0-9A-Za-z]{4})[\s\-\.\_]+([0-9A-Za-z]{4})\b",
            line,
        )
        if m_cand:
            g1 = m_cand.group(1).translate(trans_table)
            g2 = m_cand.group(2).translate(trans_table)
            g3 = m_cand.group(3).translate(trans_table)
            if g1.isdigit() and g2.isdigit() and g3.isdigit() and g1[0] in "23456789":
                return f"XXXX XXXX {g3}", f"{g1} {g2} {g3}"

        cleaned_digits = re.sub(r"[^0-9]", "", line.translate(trans_table))
        if len(cleaned_digits) == 12 and cleaned_digits[0] in "23456789":
            raw_id = f"{cleaned_digits[:4]} {cleaned_digits[4:8]} {cleaned_digits[8:]}"
            masked_id = f"XXXX XXXX {cleaned_digits[8:]}"
            return masked_id, raw_id

    pattern_continuous = r"(?<!\d)([2-9]\d{11})(?!\d)"
    match = re.search(pattern_continuous, text)
    if match:
        raw_num = match.group(1)
        raw_id = f"{raw_num[:4]} {raw_num[4:8]} {raw_num[8:]}"
        masked_id = f"XXXX XXXX {raw_num[8:]}"
        return masked_id, raw_id

    return "Not Found", "Not Found"


MONTH_NAME_TO_INT = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}


def parse_date_of_birth(text: str, lines: List[str]) -> str:
    """
    Noise-tolerant extraction of Date of Birth (DOB) or Year of Birth (YOB).
    Standardizes output to 'DD/MM/YYYY' or 'YYYY'.
    """
    if not text and not lines:
        return "Not Found"

    norm_text = re.sub(r"[\|!Il\]\[]+", "/", text)
    norm_text = re.sub(r"(\b[0-3]?\d[\/\-\.\s][0-1]?\d)[\/\-\.\s]?70(\d{2})\b", r"\1/20\2", norm_text)
    norm_text = re.sub(r"(\b[0-3]?\d[\/\-\.\s][0-1]?\d)[\/\-\.\s]?[7\/]20(\d{2})\b", r"\1/20\2", norm_text)
    norm_text = re.sub(r"(\b[0-3]?\d)1([0-1]\d)[7\/]20(\d{2})\b", r"\1/\2/20\3", norm_text)

    month_name_pattern = re.compile(
        r"(?i)\b([0-3]?\d)[\s\/\-\.]([a-zA-Z]{3,9})[\s\/\-\.]((?:19|20)\d{2})\b"
    )
    m_mn = month_name_pattern.search(norm_text)
    if m_mn:
        day_v = int(m_mn.group(1))
        m_str = m_mn.group(2).lower()[:3]
        yr_v = int(m_mn.group(3))
        if m_str in MONTH_NAME_TO_INT and 1 <= day_v <= 31 and 1900 <= yr_v <= 2026:
            return f"{day_v:02d}/{MONTH_NAME_TO_INT[m_str]:02d}/{yr_v}"

    dob_label_pattern = re.compile(
        r"(?i)(?:dob|d\.o\.b|date\s*of\s*birth|birth\s*date|जन्म\s*(?:तारीख|तिथि)|பிறந்த\s*(?:தேதி|நாள்)?|jnat[a-z@\s]*|jnaid[a-z@\s]*|jaatid[a-z@\s]*|girt\s*jnat[a-z@\s]*)[\s\:\.\-\|\/]*"
        r"([0-3]?\d)[\s\/\-\.]{1,3}([0-1]?\d)[\s\/\-\.]{1,3}((?:19|20)\d{2})\b"
    )
    m = dob_label_pattern.search(norm_text)
    if m:
        try:
            day, month, year = int(m.group(1)), int(m.group(2)), int(m.group(3))
            if 1 <= day <= 31 and 1 <= month <= 12 and 1900 <= year <= 2026:
                return f"{day:02d}/{month:02d}/{year}"
        except ValueError:
            pass

    yob_label_pattern = re.compile(
        r"(?i)(?:yob|y\.o\.b|year\s*of\s*birth|birth\s*year|जन्म\s*वर्ष|பிறந்த\s*ஆண்டு)[\s\:\.\-\|\/]*((?:19|20)\d{2})\b"
    )
    m = yob_label_pattern.search(norm_text)
    if m:
        return m.group(1)

    date_regex = re.compile(r"\b([0-3]?\d)[\s\/\-\.]{1,3}([0-1]?\d)[\s\/\-\.]{1,3}((?:19|20)\d{2})\b")
    for line in lines:
        if any(k in line.lower() for k in ["issue", "valid", "download", "print", "generation", "expiry"]):
            continue
        nl = re.sub(r"[\|!Il\]\[]+", "/", line)
        nl = re.sub(r"(\b[0-3]?\d[\/\-\.\s][0-1]?\d)[\/\-\.\s]?70(\d{2})\b", r"\1/20\2", nl)
        nl = re.sub(r"(\b[0-3]?\d[\/\-\.\s][0-1]?\d)[\/\-\.\s]?[7\/]20(\d{2})\b", r"\1/20\2", nl)
        nl = re.sub(r"(\b[0-3]?\d)1([0-1]\d)[7\/]20(\d{2})\b", r"\1/\2/20\3", nl)
        for match in date_regex.finditer(nl):
            try:
                day, month, year = int(match.group(1)), int(match.group(2)), int(match.group(3))
                if 1 <= day <= 31 and 1 <= month <= 12 and 1900 <= year <= 2026:
                    return f"{day:02d}/{month:02d}/{year}"
            except ValueError:
                continue

    return "Not Found"


def parse_gender(text: str, lines: List[str]) -> str:
    """
    Extracts Gender ('Male', 'Female', or 'Other') across English, Hindi, Tamil,
    and common OCR misreadings.
    """
    combined = f"{text}\n" + "\n".join(lines)
    if not combined.strip():
        return "Not Found"

    if re.search(r"(?i)\b(?:female|femal|feyale|cuekfeyale)\b|महिला|स्त्री|பெண்", combined):
        return "Female"
    if re.search(r"(?i)\b(?:transgender|other)\b|तृतीय\s*लिंग|திருநங்கை", combined):
        return "Other"
    if re.search(r"(?i)\b(?:male)\b|पुरुष|ஆண்", combined):
        return "Male"
    return "Not Found"


def parse_mobile_number(text: str, lines: List[str], doc_id: str = "") -> str:
    """
    Extracts a 10-digit Indian mobile number starting with 6, 7, 8, or 9.
    Supports spaced 5-5 format ('93603 34347') and hyphenated format.
    Ensures zero collision with 12-digit identity numbers or dates.
    """
    if not text and not lines:
        return "Not Found"

    sanitized = f"{text}\n" + "\n".join(lines)
    if doc_id and doc_id != "Not Found":
        raw_digits = re.sub(r"\D", "", doc_id)
        if len(raw_digits) == 12:
            sanitized = sanitized.replace(raw_digits, " ")
            sanitized = sanitized.replace(f"{raw_digits[:4]} {raw_digits[4:8]} {raw_digits[8:]}", " ")

    sanitized = re.sub(r"\b[2-9]\d{3}[\s\-\.]*\d{4}[\s\-\.]*\d{4}\b", " ", sanitized)
    sanitized = re.sub(r"\b\d{11,16}\b", " ", sanitized)

    labeled_pattern = re.compile(
        r"(?i)(?:mob(?:ile)?|mub|phone|cell|contact|tele|दूरभाष|मोबाइल)[\s\:\.\-\#]*"
        r"(?:\+?91[\s\-]*)?([6-9]\d{4}[\s\-]??\d{5}|[6-9]\d{2}[\s\-]??\d{3}[\s\-]??\d{4}|[6-9]\d{9})\b"
    )
    m = labeled_pattern.search(sanitized)
    if m:
        clean_num = re.sub(r"\D", "", m.group(1))
        if len(clean_num) == 10:
            return clean_num

    prefix_pattern = re.compile(r"(?:\+91[\s\.\-]?|(?<!\d)0)([6-9]\d{4}[\s\-]??\d{5}|[6-9]\d{9})(?!\d)")
    m = prefix_pattern.search(sanitized)
    if m:
        clean_num = re.sub(r"\D", "", m.group(1))
        if len(clean_num) == 10:
            return clean_num

    standalone_pattern = re.compile(r"(?<!\d)([6-9]\d{4}\s\d{5})(?!\d)|(?<!\d)([6-9]\d{9})(?!\d)")
    for match in standalone_pattern.finditer(sanitized):
        val = match.group(1) or match.group(2)
        if val:
            clean = re.sub(r"\D", "", val)
            if len(clean) == 10:
                return clean

    return "Not Found"


def clean_name_string(s: str) -> str:
    """Cleans whitespace and punctuation artifacts in extracted Indian names."""
    s = re.sub(r"[^A-Za-z\s\.]", " ", s)
    s = re.sub(r"\s+", " ", s).strip(" .")
    return s


def parse_father_or_guardian_name(lines: List[str]) -> str:
    """
    Extracts Father, Guardian, or Spouse Name from OCR lines or Address blocks.
    Supports C/O, S/O, D/O, W/O, D'o, C.O, S.O, Care of, Father's Name, and OCR variants.
    Captures multi-line continuations and strips downstream structural address tokens.
    """
    trigger_regex = re.compile(
        r"(?i)\b(?:"
        r"[CSDWcsdw][\/\'\.\-\\][0oO]|"
        r"(?:C[0oO]|S[0oO]|D[0oO]|W[0oO]|CQ|CIO|SIO|DIO|WIO|[5$sSdDcC][\/\'\.\-\\][Qq0oO])\b[\s\:\.\-\,\/]*|"
        r"(?:C[oO]|S[oO]|D[oO])(?=[A-Z][a-z]+)|"
        r"Father(?:\'?s)?\s*Name|Care\s*of|Son\s*of|Daughter\s*of|Wife\s*of|Guardian(?:\'?s)?\s*Name|"
        r"पिता|पति|सुपुत्र|आत्मज|தந்தை|க\/பெ|த\/பெ"
        r")[\s\:\.\-\,\/]*"
    )

    address_cutoffs = [
        "address", "पता", "house no", "house", "h.no", "h no", "h-no", "door no", "door", "flat",
        "plot", "street", "road", "rd", "cross", "main", "lane", "gali",
        "layout", "nagar", "colony", "enclave", "village", "post office", "post", "p.o.", "p.o",
        "dist", "district", "taluk", "tehsil", "mandal", "pin",
        "pincode", "near", "opp", "opposite", "behind", "beside", "floor",
        "bldg", "building", "apartment", "apt", "block", "sector", "phase",
    ]

    system_stop_words = {
        "government", "india", "unique", "identification", "authority",
        "aadhaar", "card", "enrolment", "male", "female", "help", "www",
        "uidai", "gov", "in", "date", "birth", "year", "address", "floor",
        "house", "road", "street", "nagar", "village", "pin",
    }

    candidates = []
    for idx, line in enumerate(lines):
        clean_line = line.strip()
        if "%" in clean_line or "#" in clean_line or "*" in clean_line:
            continue

        m = trigger_regex.search(clean_line)
        if m:
            remainder = clean_line[m.end():].strip()
            candidate = ""
            if len(remainder.split()) >= 1 and re.search(r"[A-Za-z]", remainder):
                candidate = remainder
            elif idx + 1 < len(lines):
                candidate = lines[idx + 1].strip()

            if not candidate:
                continue

            for cutoff in address_cutoffs:
                cut_match = re.search(rf"\b{re.escape(cutoff)}\b", candidate, re.IGNORECASE)
                if cut_match:
                    candidate = candidate[:cut_match.start()].strip()

            candidate = re.split(r"[,;:\-–—\d]", candidate)[0].strip()
            candidate = clean_name_string(candidate)

            words = candidate.split()
            if 1 <= len(words) <= 4:
                if not any(w.lower() in system_stop_words for w in words):
                    valid_words = [w for w in words if len(w) > 1 or w.isupper()]
                    if valid_words:
                        candidates.append(" ".join(valid_words).title())

    if candidates:
        counts = collections.Counter(candidates)
        return counts.most_common(1)[0][0]

    return "Not Found"


def score_holder_name_candidate(line: str, is_near_dob: bool = False) -> Tuple[int, str]:
    """Scores a candidate line for likelihood of being the document holder name."""
    if any(c in line for c in ["%", "@", "#", "*", "$", "(", ")", "[", "]", "{", "}", "/", "\\", ":", ";", "=", "_"]):
        return 0, ""

    cand_clean = clean_name_string(line)
    words = cand_clean.split()
    if not (1 <= len(words) <= 4):
        return 0, ""

    system_stop_words = {
        "government", "india", "unique", "identification", "authority",
        "aadhaar", "card", "enrolment", "male", "female", "help", "www",
        "uidai", "gov", "in", "date", "birth", "year", "dob", "yob",
        "address", "father", "mother", "guardian", "husband", "mob",
        "mub", "mobile", "email", "mail", "phone", "contact", "vid",
        "information", "your", "no", "info", "grdine", "eeim", "jnatd",
        "jnaid", "jaatid", "gxl", "additt", "glog", "girt", "cuekfeyale",
        "pirantha", "thethi", "naal", "signature", "photo", "valid",
        "republic", "state", "national", "to",
    }

    if any(w.lower() in system_stop_words for w in words):
        return 0, ""

    if not all(len(w) >= 2 or (len(w) == 1 and w.isupper()) for w in words):
        return 0, ""

    score = 0
    if len(words) == 2:
        score += 35
    elif len(words) == 3:
        score += 30
    elif len(words) == 1:
        if len(words[0]) >= 4 and any(c in "aeiouAEIOU" for c in words[0]):
            score += 5
        else:
            return 0, ""

    if cand_clean.isupper():
        score += 15
    elif cand_clean.istitle():
        score += 12
    else:
        if sum(1 for w in words if w[0].isupper()) >= len(words) / 2:
            score += 10

    if is_near_dob:
        score += 25

    return score, cand_clean.title()


def parse_holder_name(lines: List[str]) -> str:
    """
    Extracts the document holder's name using multi-factor likelihood scoring.
    """
    dob_idx = -1
    for idx, l in enumerate(lines):
        if any(k in l.lower() for k in ["dob", "date of birth", "birth", "yob", "female", "male", "jnat", "jnaid", "jaatid"]):
            dob_idx = idx
            break

    candidates: List[Tuple[int, int, str]] = []
    for idx, line in enumerate(lines[:35]):
        is_near = dob_idx != -1 and 0 < (dob_idx - idx) <= 4
        score, name = score_holder_name_candidate(line, is_near_dob=is_near)
        if score > 0:
            candidates.append((score, idx, name))

    if candidates:
        candidates.sort(
            key=lambda x: (x[0], -abs((dob_idx if dob_idx != -1 else 0) - x[1])),
            reverse=True,
        )
        return candidates[0][2]

    return "Not Found"


def parse_address(text: str, lines: List[str]) -> str:
    """
    Extracts residential address from OCR lines or text blocks, capturing from 'Address:'
    or relationship markers up to the 6-digit Indian PIN code.
    """
    if not lines and text:
        lines = [l.strip() for l in text.splitlines() if l.strip()]
    if not lines:
        return "Not Found"

    addr_start_idx = -1
    for idx, line in enumerate(lines):
        if re.search(r"(?i)\b(?:address|पता|முகவரி)\b", line):
            addr_start_idx = idx
            break

    if addr_start_idx == -1:
        for idx, line in enumerate(lines):
            if re.search(r"(?i)\b(?:[CSDW][\/\.][O0]|Care\s*of|House\s*No|Door\s*No|Flat|Street|Road|Village|PIN)\b", line):
                addr_start_idx = idx
                break

    if addr_start_idx == -1:
        return "Not Found"

    collected: List[str] = []
    for line in lines[addr_start_idx:addr_start_idx + 7]:
        clean_l = re.sub(r"(?i)^(?:address|पता|முகவரி)[\s\:\-\.]*", "", line).strip()
        if not clean_l:
            continue
        if any(k in clean_l.lower() for k in ["help@", "www.uidai", "1947", "dob:", "date of birth"]):
            break
        if re.match(r"^[2-9]\d{3}\s\d{4}\s\d{4}$", clean_l):
            break
        collected.append(clean_l)
        if re.search(r"\b[1-9]\d{5}\b", clean_l):
            break

    if collected:
        return ", ".join(collected)
    return "Not Found"


def hybrid_fuse_fields(
    vlm_data: Optional[Dict[str, Any]],
    ocr_text: str,
    ocr_lines: List[str],
    privacy_mask: bool = True,
) -> Dict[str, str]:
    """
    3-Stage Hybrid Aadhaar Field Fusion:
    Combines Qwen2-VL-2B structured JSON output with deterministic regex parsers
    and fallback OCR / Digital PDF text lines, including cross-field extraction
    of Father/Guardian Name and Mobile Number from Address blocks.
    """
    vlm = vlm_data or {}
    vlm_id_str = str(vlm.get("document_id", "Not Found")).strip()
    vlm_holder_str = str(vlm.get("holder_name", "Not Found")).strip()
    vlm_dob_str = str(vlm.get("date_of_birth", "Not Found")).strip()
    vlm_gender_str = str(vlm.get("gender", "Not Found")).strip()
    vlm_father_str = str(vlm.get("father_or_guardian", "Not Found")).strip()
    vlm_mobile_str = str(vlm.get("mobile_number", "Not Found")).strip()
    vlm_addr_str = str(vlm.get("address", "Not Found")).strip()

    # 1. Document ID (VLM -> Regex Normalizer -> OCR Fallback)
    masked_id, raw_id = parse_document_id(vlm_id_str, [vlm_id_str])
    if raw_id == "Not Found":
        masked_id, raw_id = parse_document_id(ocr_text, ocr_lines)
    display_id = masked_id if privacy_mask else raw_id

    # 2. Date of Birth (VLM -> Regex Normalizer -> OCR Fallback)
    final_dob = parse_date_of_birth(vlm_dob_str, [vlm_dob_str])
    if final_dob == "Not Found":
        final_dob = parse_date_of_birth(ocr_text, ocr_lines)

    # 3. Gender (VLM -> OCR Fallback)
    final_gender = parse_gender(vlm_gender_str, [vlm_gender_str])
    if final_gender == "Not Found":
        final_gender = parse_gender(ocr_text, ocr_lines)

    # 4. Address (VLM -> OCR Fallback)
    if vlm_addr_str and vlm_addr_str != "Not Found":
        final_address = vlm_addr_str
    else:
        final_address = parse_address(ocr_text, ocr_lines)

    # 5. Father / Guardian Name (VLM -> VLM Address Cross-Check -> OCR Lines Fallback)
    final_father = "Not Found"
    if vlm_father_str and vlm_father_str != "Not Found":
        parsed_from_vlm = parse_father_or_guardian_name([f"C/O: {vlm_father_str}"])
        if parsed_from_vlm != "Not Found":
            final_father = parsed_from_vlm
        else:
            cleaned_f = re.sub(
                r"(?i)^(?:[csdw][/'\.\s_\-]?[o0]|father['\s]*s\s*name|care\s*of)[\s:\.\-=,]*",
                "",
                vlm_father_str,
            ).strip()
            cleaned_f = clean_name_string(cleaned_f)
            if cleaned_f:
                final_father = cleaned_f.title()

    if final_father == "Not Found" and vlm_addr_str and vlm_addr_str != "Not Found":
        final_father = parse_father_or_guardian_name([vlm_addr_str] + vlm_addr_str.splitlines())

    if final_father == "Not Found" and ocr_lines:
        final_father = parse_father_or_guardian_name(ocr_lines)

    # 6. Mobile Number (VLM -> VLM Address -> OCR Corpus, preventing 12-digit ID collision)
    final_mobile = parse_mobile_number(vlm_mobile_str, [vlm_mobile_str], doc_id=raw_id)
    if final_mobile == "Not Found" and vlm_addr_str != "Not Found":
        final_mobile = parse_mobile_number(vlm_addr_str, [vlm_addr_str], doc_id=raw_id)
    if final_mobile == "Not Found":
        final_mobile = parse_mobile_number(ocr_text, ocr_lines, doc_id=raw_id)

    # 7. Holder Name (VLM -> OCR Scoring Fallback)
    final_holder = "Not Found"
    if vlm_holder_str and vlm_holder_str != "Not Found":
        cleaned_h = clean_name_string(vlm_holder_str)
        if cleaned_h:
            final_holder = cleaned_h.title()
    if final_holder == "Not Found" and ocr_lines:
        final_holder = parse_holder_name(ocr_lines)

    return {
        "Document ID": display_id,
        "Raw Document ID": raw_id,
        "Holder Name": final_holder,
        "Date of Birth": final_dob,
        "Gender": final_gender,
        "Father / Guardian Name": final_father,
        "Mobile Number": final_mobile,
        "Address": final_address,
    }


def merge_page_records(page_records: List[Dict[str, Any]]) -> Dict[str, Any]:
    """
    Merges extracted field dictionaries across multiple pages (e.g. Page 1 Front + Page 2 Back)
    so fields located on the back of an Aadhaar card (Father/Guardian Name, Address, Mobile)
    seamlessly combine with front-side fields (Holder Name, DOB, Gender, Document ID).
    """
    fields = [
        "Document ID",
        "Raw Document ID",
        "Holder Name",
        "Date of Birth",
        "Gender",
        "Father / Guardian Name",
        "Mobile Number",
        "Address",
    ]
    merged: Dict[str, Any] = {f: "Not Found" for f in fields}

    for rec in page_records:
        for f in fields:
            if merged[f] == "Not Found":
                val = str(rec.get(f, "Not Found")).strip()
                if val and val != "Not Found":
                    merged[f] = val

    return merged


# ==============================================================================
# EXCEL EXPORT ENGINE VIA OPENPYXL
# ==============================================================================

def export_records_to_excel(records: List[Dict[str, Any]]) -> bytes:
    """
    Exports structured extraction records to a professionally styled Excel (.xlsx) file.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Extracted ID Records"

    headers = [
        "File Target",
        "Document ID",
        "Holder Name",
        "Date of Birth",
        "Gender",
        "Father / Guardian Name",
        "Mobile Number",
        "Address",
        "Extraction Mode",
        "Status",
    ]
    ws.append(headers)

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )
    data_font = Font(name="Segoe UI", size=10)
    data_align = Alignment(vertical="center")

    for row_idx, rec in enumerate(records, start=2):
        row_vals = [rec.get(h, "Not Found") for h in headers]
        ws.append(row_vals)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = data_align

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(min(max_len + 4, 50), 15)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ==============================================================================
# STREAMLIT APPLICATION INTERFACE
# ==============================================================================

def main():
    st.set_page_config(
        page_title="Hybrid Aadhaar & ID Document Extractor",
        page_icon="🪪",
        layout="wide",
        initial_sidebar_state="expanded",
    )

    st.markdown(
        """
    <style>
        .main-header {
            font-size: 2.1rem;
            font-weight: 800;
            color: #1e293b;
            margin-bottom: 0.2rem;
        }
        .sub-header {
            font-size: 1.0rem;
            color: #64748b;
            margin-bottom: 1.5rem;
        }
        .field-card {
            background: linear-gradient(135deg, #ffffff 0%, #f8fafc 100%);
            border: 1px solid #e2e8f0;
            border-radius: 12px;
            padding: 16px 18px;
            box-shadow: 0 2px 4px rgba(0, 0, 0, 0.04);
            margin-bottom: 12px;
            transition: transform 0.15s ease, box-shadow 0.15s ease;
        }
        .field-card:hover {
            box-shadow: 0 4px 10px rgba(0, 0, 0, 0.08);
            transform: translateY(-2px);
        }
        .field-label {
            font-size: 0.78rem;
            font-weight: 700;
            text-transform: uppercase;
            letter-spacing: 0.6px;
            color: #64748b;
            margin-bottom: 4px;
            display: flex;
            align-items: center;
            gap: 6px;
        }
        .field-value {
            font-size: 1.15rem;
            font-weight: 700;
            color: #0f172a;
            word-break: break-word;
            font-family: 'Consolas', 'Courier New', monospace;
        }
        .field-badge {
            display: inline-block;
            font-size: 0.7rem;
            font-weight: 600;
            padding: 2px 8px;
            border-radius: 9999px;
            margin-left: 8px;
        }
        .badge-success {
            background-color: #dcfce7;
            color: #166534;
        }
        .badge-missing {
            background-color: #fee2e2;
            color: #991b1b;
        }
    </style>
    """,
        unsafe_allow_html=True,
    )

    if "results" not in st.session_state:
        st.session_state.results = []
    if "latest_debug" not in st.session_state:
        st.session_state.latest_debug = {}

    st.markdown('<div class="main-header">🪪 Hybrid Aadhaar & Identity Document Extractor</div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="sub-header">Multi-Modal Vision-Language (Qwen2-VL-2B) + Dual-Pass EasyOCR + PyMuPDF Pipeline. '
        "Extracts Aadhaar ID, Holder Name, DOB, Gender, Father/Guardian Name, Mobile Number, and Address.</div>",
        unsafe_allow_html=True,
    )

    with st.sidebar:
        st.header("⚙️ Engine Configuration")
        engine_choice = st.selectbox(
            "🧠 Extraction Engine",
            [
                "Hybrid (Qwen2-VL-2B + EasyOCR + Regex)",
                "Fast Local OCR (EasyOCR + PyMuPDF + Regex)",
            ],
            index=0 if QWEN_VL_AVAILABLE else 1,
        )
        custom_qwen_path = st.text_input(
            "📂 Qwen2-VL-2B Drive / Local Path",
            value="/content/drive/MyDrive/model_ocr/qwen2_vl_2b",
            help="Path to your Qwen2-VL-2B folder on Google Drive or local disk. Automatically verified and repaired if incomplete.",
        )
        auto_repair_toggle = st.checkbox(
            "🛠️ Auto-Repair Missing Drive Shards",
            value=True,
            help="Automatically downloads missing .safetensors shards from Hugging Face if the Drive folder is incomplete.",
        )
        mask_id_toggle = st.checkbox(
            "🔒 Mask Document ID (Privacy)",
            value=True,
            help="Hides first 8 digits as 'XXXX XXXX 1234'",
        )
        use_clahe_toggle = st.checkbox(
            "✨ Dual-Pass CLAHE Preprocessing",
            value=True,
            help="Aggregates standard grayscale + mild CLAHE for faint text",
        )
        ocr_lang_option = st.selectbox(
            "🌐 OCR Language Models",
            ["English ('en')", "English + Tamil ('en', 'ta')", "English + Hindi ('en', 'hi')"],
            index=0,
        )
        if "Tamil" in ocr_lang_option:
            languages = ("en", "ta")
        elif "Hindi" in ocr_lang_option:
            languages = ("en", "hi")
        else:
            languages = ("en",)

    reader = get_easyocr_reader(languages)
    use_qwen = "Qwen2-VL" in engine_choice and QWEN_VL_AVAILABLE
    qwen_model, qwen_processor, resolved_qwen_path = (None, None, "Disabled")
    if use_qwen:
        qwen_model, qwen_processor, resolved_qwen_path = load_qwen_vl_model(
            custom_path=custom_qwen_path.strip() or None,
            auto_repair=auto_repair_toggle,
        )

    tab_gdrive, tab_upload = st.tabs(["🔗 Google Drive Link", "📁 Direct File Upload"])

    gdrive_link = ""
    with tab_gdrive:
        st.markdown("Enter a public Google Drive link (file, uc, or folder URL) or raw File ID:")
        gdrive_link = st.text_input(
            "Google Drive URL / ID:",
            placeholder="https://drive.google.com/file/d/1A2B3C4D.../view?usp=sharing",
        )

    uploaded_files = []
    with tab_upload:
        uploaded_files = st.file_uploader(
            "Upload Identity Documents (Single/Multi-Page PDFs or Images):",
            type=["pdf", "png", "jpg", "jpeg", "tiff", "bmp", "webp"],
            accept_multiple_files=True,
        )

    col_btn, _ = st.columns([1, 4])
    with col_btn:
        start_processing = st.button("🚀 Process Document", type="primary", use_container_width=True)

    if start_processing:
        if not gdrive_link and not uploaded_files:
            st.error("⚠️ Please provide a Google Drive link or upload at least one document.")
            return

        with st.spinner("Processing document through Hybrid Aadhaar Vision Pipeline..."):
            with tempfile.TemporaryDirectory() as temp_dir:
                files_to_process: List[str] = []

                if gdrive_link:
                    try:
                        st.info("Fetching document from Google Drive...")
                        downloaded_paths = download_from_gdrive(gdrive_link, temp_dir)
                        files_to_process.extend(downloaded_paths)
                    except Exception as e:
                        st.error(f"❌ Google Drive Download Failed: {e}")
                        return

                if uploaded_files:
                    for uf in uploaded_files:
                        p = os.path.join(temp_dir, uf.name)
                        with open(p, "wb") as f:
                            f.write(uf.getbuffer())
                        files_to_process.append(p)

                for fpath in files_to_process:
                    base_name = os.path.basename(fpath)
                    start_time = time.time()
                    extraction_mode = (
                        f"Hybrid Qwen2-VL ({os.path.basename(resolved_qwen_path)}) + OCR"
                        if (use_qwen and qwen_model is not None)
                        else "OCR Dual-Pass (Grayscale + CLAHE)"
                    )

                    dig_text, dig_lines = ("", [])
                    if fpath.lower().endswith(".pdf"):
                        dig_text, dig_lines = extract_pdf_digital_text(fpath)

                    try:
                        pages = load_document_pages(fpath, base_name)
                    except Exception as e:
                        st.error(f"Error loading document pages for {base_name}: {e}")
                        continue

                    page_records: List[Dict[str, Any]] = []
                    all_ocr_lines: List[str] = list(dig_lines)
                    vlm_raw_outputs: List[str] = []

                    for page_img, page_title in pages:
                        p_ocr_text, p_ocr_lines = run_dual_pass_ocr(
                            page_img, reader, use_clahe_pass=use_clahe_toggle
                        )
                        combined_lines = list(dig_lines) + p_ocr_lines
                        combined_text = "\n".join(combined_lines)
                        all_ocr_lines.extend(p_ocr_lines)

                        vlm_res = None
                        if use_qwen and qwen_model is not None:
                            pil_page = Image.fromarray(cv2.cvtColor(page_img, cv2.COLOR_BGR2RGB))
                            vlm_res = extract_fields_with_qwen(pil_page, qwen_model, qwen_processor)
                            if vlm_res.get("raw_output"):
                                vlm_raw_outputs.append(f"[{page_title}] {vlm_res['raw_output']}")

                        fused_page = hybrid_fuse_fields(
                            vlm_data=vlm_res,
                            ocr_text=combined_text,
                            ocr_lines=combined_lines,
                            privacy_mask=mask_id_toggle,
                        )
                        page_records.append(fused_page)

                    merged = merge_page_records(page_records)
                    proc_time = round(time.time() - start_time, 2)
                    status_flag = (
                        "Success"
                        if (merged["Document ID"] != "Not Found" and merged["Date of Birth"] != "Not Found")
                        else "Partial"
                    )

                    record = {
                        "File Target": base_name,
                        "Document ID": merged["Document ID"],
                        "Holder Name": merged["Holder Name"],
                        "Date of Birth": merged["Date of Birth"],
                        "Gender": merged["Gender"],
                        "Father / Guardian Name": merged["Father / Guardian Name"],
                        "Mobile Number": merged["Mobile Number"],
                        "Address": merged["Address"],
                        "Extraction Mode": extraction_mode,
                        "Status": status_flag,
                    }
                    st.session_state.results.append(record)

                    st.session_state.latest_debug = {
                        "file_name": base_name,
                        "lines": all_ocr_lines,
                        "corpus": "\n".join(all_ocr_lines),
                        "vlm_outputs": "\n\n".join(vlm_raw_outputs),
                        "time_sec": proc_time,
                        "mode": extraction_mode,
                    }

        st.success(f"Processing complete in {st.session_state.latest_debug.get('time_sec', 0)} seconds!")

    if st.session_state.results:
        latest = st.session_state.results[-1]
        st.markdown("### 📋 Extracted Field Summary (Latest Document)")

        r1_cols = st.columns(4)
        top_fields = [
            ("🪪 12-Digit Document ID", "Document ID"),
            ("🏷️ Holder Name", "Holder Name"),
            ("🎂 Date of Birth", "Date of Birth"),
            ("⚧️ Gender", "Gender"),
        ]
        for col, (label, key) in zip(r1_cols, top_fields):
            with col:
                val = latest.get(key, "Not Found")
                badge = "badge-success" if val != "Not Found" else "badge-missing"
                st.markdown(
                    f"""
                <div class="field-card">
                    <div class="field-label">{label} <span class="field-badge {badge}">{'Found' if val != 'Not Found' else 'Missing'}</span></div>
                    <div class="field-value">{val}</div>
                </div>
                """,
                    unsafe_allow_html=True,
                )

        r2_cols = st.columns([1, 1, 2])
        bottom_fields = [
            ("👤 Father / Guardian", "Father / Guardian Name"),
            ("📱 Mobile Number", "Mobile Number"),
            ("🏠 Residential Address", "Address"),
        ]
        for col, (label, key) in zip(r2_cols, bottom_fields):
            with col:
                val = latest.get(key, "Not Found")
                badge = "badge-success" if val != "Not Found" else "badge-missing"
                st.markdown(
                    f"""
                <div class="field-card">
                    <div class="field-label">{label} <span class="field-badge {badge}">{'Found' if val != 'Not Found' else 'Missing'}</span></div>
                    <div class="field-value">{val}</div>
                </div>
                """,
                    unsafe_allow_html=True,
                )

        st.caption(
            f"Source: `{latest['File Target']}` | Mode: `{latest['Extraction Mode']}` | Status: `{latest['Status']}`"
        )

        st.markdown("### 📊 Extracted Records Database")
        df_results = pd.DataFrame(st.session_state.results)
        edited_df = st.data_editor(
            df_results,
            use_container_width=True,
            num_rows="dynamic",
            key="data_editor_widget",
        )

        col_export_xlsx, col_export_csv, col_clear = st.columns([2, 2, 6])
        with col_export_xlsx:
            try:
                excel_bytes = export_records_to_excel(edited_df.to_dict("records"))
                st.download_button(
                    label="📥 Export to Excel (.xlsx)",
                    data=excel_bytes,
                    file_name="ID_Extracted_Records.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                )
            except Exception as e:
                st.error(f"Excel export error: {e}")

        with col_export_csv:
            csv_bytes = edited_df.to_csv(index=False).encode("utf-8")
            st.download_button(
                label="📄 Export to CSV (.csv)",
                data=csv_bytes,
                file_name="ID_Extracted_Records.csv",
                mime="text/csv",
                use_container_width=True,
            )

        with col_clear:
            if st.button("🗑️ Clear All Records"):
                st.session_state.results = []
                st.session_state.latest_debug = {}
                st.rerun()

    if st.session_state.latest_debug:
        with st.expander("🔍 Raw VLM / OCR & Pipeline Debug Inspector", expanded=False):
            dbg = st.session_state.latest_debug
            st.write(f"**Target File:** `{dbg.get('file_name', 'N/A')}`")
            st.write(f"**Pipeline Mode:** `{dbg.get('mode', 'N/A')}`")
            st.write(f"**Execution Duration:** `{dbg.get('time_sec', 0)} seconds`")
            if dbg.get("vlm_outputs"):
                st.markdown("**Qwen2-VL-2B Raw JSON Output:**")
                st.code(dbg["vlm_outputs"], language="json")
            col_lines, col_raw = st.columns(2)
            with col_lines:
                st.markdown("**Processed Lines (Reading Order):**")
                for idx, line in enumerate(dbg.get("lines", []), start=1):
                    st.text(f"[{idx:02d}] {line}")
            with col_raw:
                st.markdown("**Aggregated Raw Text Corpus:**")
                st.code(dbg.get("corpus", "No text extracted"), language="text")


if __name__ == "__main__":
    main()
